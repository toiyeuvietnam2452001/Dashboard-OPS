"""Add dashboard sheets (formulas + native charts) in front of the data tables in OPS_Dashboard_Data.xlsx.

usage: excel_dash.py <in.xlsx> <out.xlsx>
Dashboard sheets hold no Excel Tables, so the Office Script that exports tables for the web dashboard ignores them.
"""
import sys
from openpyxl import load_workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.axis import TextAxis
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import DataBarRule, FormulaRule, CellIsRule
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

SRC, OUT = sys.argv[1:3]
wb = load_workbook(SRC)
N = 300  # rows covered by every data range, leaves room for new rows

# ---- column references into the data tables -------------------------------------------------
TAB = {}
for ws in wb.worksheets:
    for name, ref in ws.tables.items():
        head = [c.value for c in ws[ref][0]]
        TAB[name] = (ws.title, {h: L(i + 1) for i, h in enumerate(head)})

def R(table, col, absolute=True):
    sheet, cols = TAB[table]
    c = cols[col]
    return f"'{sheet}'!${c}$2:${c}${N}"

def ROWS(table):  # for ROW() based lookups
    sheet, cols = TAB[table]
    return f"'{sheet}'!$A$2:$A${N}"

# ---- styles ----------------------------------------------------------------------------------
NAVY, BLUE, RED, AMBER, GREEN, GREY, MUTED, LINE, BG = '1B2430', '1F6FB2', 'D1495B', 'E0A526', '1F7A4D', '9AA5B4', '5D6878', 'E1E5EB', 'F3F5F8'
WHITE = PatternFill('solid', fgColor='FFFFFF'); PAGE = PatternFill('solid', fgColor=BG)
HEADF = PatternFill('solid', fgColor='E8EEF5')
thin = Side(style='thin', color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=thin)
F = lambda **k: Font(name='Calibri', **{'size': 10, 'color': NAVY, **k})

def page(title, tab_color):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    for c in range(1, 27):
        ws.column_dimensions[L(c)].width = 11
    ws.column_dimensions['A'].width = 2
    for r in range(1, 140):
        for c in range(1, 27):
            ws.cell(r, c).fill = PAGE
    return ws

def header(ws, subtitle):
    ws['B1'] = '="OPS readiness – report "&TEXT(DATEVALUE(\'Thong tin\'!$A$2),"d mmm yyyy")'
    ws['B1'].font = F(size=20)
    ws['B2'] = subtitle
    ws['B2'].font = F(size=10, color=MUTED, italic=True)
    ws.row_dimensions[1].height = 30

def card(ws, r, c, w, label, value, fmt='General', note=None, value_color=NAVY):
    """KPI tile: label / big value / note, w columns wide starting at (r, c)."""
    for i, (text, font, h) in enumerate([(label, F(size=9, color=MUTED), 16), (value, F(size=22, color=value_color), 34), (note, F(size=9, color=MUTED), 30)]):
        ws.merge_cells(start_row=r + i, start_column=c, end_row=r + i, end_column=c + w - 1)
        cell = ws.cell(r + i, c, text)
        cell.font = font
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
        if i == 1: cell.number_format = fmt
        ws.row_dimensions[r + i].height = max(ws.row_dimensions[r + i].height or 0, h)
        for cc in range(c, c + w):
            ws.cell(r + i, cc).fill = WHITE
            ws.cell(r + i, cc).border = Border(left=thin if cc == c else None, right=thin if cc == c + w - 1 else None,
                                               top=thin if i == 0 else None, bottom=thin if i == 2 else None)

def section(ws, r, c, text, note=None):
    ws.cell(r, c, text).font = F(size=13, bold=True)
    if note:
        ws.cell(r + 1, c, note).font = F(size=9, color=MUTED, italic=True)

def table(ws, r, c, heads, widths=None):
    for i, h in enumerate(heads):
        cell = ws.cell(r, c + i, h)
        cell.font = F(size=9, bold=True, color=MUTED); cell.fill = HEADF; cell.border = BOTTOM
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    if widths:
        for i, wd in enumerate(widths):
            if wd: ws.column_dimensions[L(c + i)].width = wd

def body(ws, r, c, formulas, n, fmts=None, wrap=None):
    """n rows of formulas; each formula is a function of k (1-based row in the source)."""
    for k in range(1, n + 1):
        for i, f in enumerate(formulas):
            cell = ws.cell(r + k - 1, c + i, f(k))
            cell.font = F(); cell.fill = WHITE; cell.border = BOTTOM
            cell.alignment = Alignment(vertical='top', wrap_text=bool(wrap and wrap[i]))
            if fmts and fmts[i]: cell.number_format = fmts[i]

def pick(table, col, k):  # k-th data row, blank when empty
    return f'=IF(INDEX({R(table, col)},{k})="","",INDEX({R(table, col)},{k}))'

def nth_where(table, cond, k):
    """row number (1-based) of the k-th data row matching cond (an array expression), '' when none."""
    return f'=IFERROR(SMALL(IF({cond},ROW({ROWS(table)})-1),{k}),"")'

def by_idx(table, col, idx_cell):
    return f'=IF({idx_cell}="","",IF(INDEX({R(table, col)},{idx_cell})="","",INDEX({R(table, col)},{idx_cell})))'

def bar(ws, rng, color=BLUE):
    ws.conditional_formatting.add(rng, DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=color, showValue=True))

def red_below(ws, rng, first, limit=0.5):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({first}<>"",{first}<{limit})'], font=Font(color=RED, bold=True)))

def line_chart(title, ytitle, w=24, h=8):
    ch = LineChart(); ch.title = title; ch.y_axis.title = ytitle; ch.height = h; ch.width = w
    ch.legend.position = 'b'; ch.y_axis.majorGridlines = None
    ch.x_axis = TextAxis(); ch.x_axis.auto = False; ch.x_axis.number_format = '@'
    ch.y_axis.delete = False; ch.x_axis.delete = False
    return ch

def style_series(s, color, dashed=False, marker=False):
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = 22000
    if dashed: s.graphicalProperties.line.dashStyle = 'dash'
    s.smooth = False
    if marker:
        s.marker.symbol = 'circle'; s.marker.size = 6
        s.marker.graphicalProperties = GraphicalProperties(solidFill=color); s.marker.graphicalProperties.line.solidFill = color
    else:
        s.marker.symbol = 'none'

def data_ref(table, col, n):
    sheet, cols = TAB[table]
    ci = ord(cols[col]) - 64
    return Reference(wb[sheet], min_col=ci, min_row=1, max_row=n + 1)

def cat_ref(table, col, n, col2=None):
    sheet, cols = TAB[table]
    ci = ord(cols[col]) - 64
    return Reference(wb[sheet], min_col=ci, max_col=(ord(cols[col2]) - 64) if col2 else ci, min_row=2, max_row=n + 1)

def n_rows(table):  # data rows only; blank rows pre-formatted for future entries don't count
    sheet, cols = TAB[table]
    ws = wb[sheet]
    return max((r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value not in (None, '')), default=1) - 1

PCT, PCT0, NUM = '0.0%', '0%', '#,##0'
LAST = lambda tbl, col: f'LOOKUP(2,1/({R(tbl, col)}<>""),{R(tbl, col)})'

# ================================================================================================
# Overview
# ================================================================================================
ov = page('Overview', BLUE)
header(ov, 'Live view of the data tables in this workbook. Edit the numbers in the data sheets (after "Huong dan"); this page recalculates at once.')
mp_month = LAST('manpower', 'month').replace('LOOKUP(2,1/(' + R('manpower', 'month') + '<>""),', 'LOOKUP(2,1/(' + R('manpower', 'total_actual') + '<>""),')
orr_month = f'LOOKUP(2,1/({R("orr_curve", "actual_pct")}<>""),{R("orr_curve", "month")})'
ov['B3'] = (f'="Manpower in "&TEXT(DATEVALUE({mp_month}&"-01"),"mmm yyyy")&": "&{LAST("manpower", "total_actual")}&" / "'
            f'&LOOKUP(2,1/({R("manpower", "total_actual")}<>""),{R("manpower", "total_plan")})&" people against plan; OPS teams ORR checklist at "'
            f'&TEXT({LAST("orr_curve", "actual_pct")}/100,"0.0%")&" vs plan "&TEXT(LOOKUP(2,1/({R("orr_curve", "actual_pct")}<>""),{R("orr_curve", "plan_pct")})/100,"0.0%")'
            f'&"; "&SUMIFS({R("pr_pipeline", "count")},{R("pr_pipeline", "status")},"Late")&" PRs late."')
ov['B3'].font = F(size=11)
import os
from openpyxl.drawing.image import Image as XLImage
logo = XLImage(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'logo.png'))
logo.height, logo.width = 62, round(62 * logo.width / logo.height)
ov.add_image(logo, 'X1')
ov.merge_cells('B3:Y3')

r = 5
mp_act, mp_plan = LAST('manpower', 'total_actual'), f'LOOKUP(2,1/({R("manpower", "total_actual")}<>""),{R("manpower", "total_plan")})'
orr_act, orr_plan = LAST('orr_curve', 'actual_pct'), f'LOOKUP(2,1/({R("orr_curve", "actual_pct")}<>""),{R("orr_curve", "plan_pct")})'
cbta = lambda p: f'SUMIFS({R("cbta", "completed")},{R("cbta", "platform")},"{p}")/SUMIFS({R("cbta", "modules")},{R("cbta", "platform")},"{p}")'
card(ov, r, 2, 4, '="OPS manpower · "&TEXT(DATEVALUE(' + mp_month + '&"-01"),"mmm yyyy")', f'={mp_act}', NUM,
     f'=TEXT({mp_act}/{mp_plan},"0%")&" of monthly plan ("&{mp_plan}&")"')
card(ov, r, 6, 4, '="OPS teams ORR checklist · "&TEXT(DATEVALUE(' + orr_month + '&"-01"),"mmm yyyy")', f'={orr_act}/100', PCT,
     f'="Plan "&TEXT({orr_plan}/100,"0.0%")&" · "&IF({orr_act}>={orr_plan},"+","")&TEXT({orr_act}-{orr_plan},"0.0")&" pts"')
card(ov, r, 10, 4, 'CBTA development (CPP & WHP)', f'=AVERAGE({cbta("CPP")},{cbta("WHP")})', PCT,
     f'="Theory assessment "&TEXT(AVERAGE(AVERAGEIFS({R("cbta", "theory_pct")},{R("cbta", "platform")},"CPP"),AVERAGEIFS({R("cbta", "theory_pct")},{R("cbta", "platform")},"WHP"))/100,"0.0%")'
     f'&" · OJT "&TEXT(AVERAGE(AVERAGEIFS({R("cbta", "ojt_pct")},{R("cbta", "platform")},"CPP"),AVERAGEIFS({R("cbta", "ojt_pct")},{R("cbta", "platform")},"WHP"))/100,"0.0%")')
card(ov, r, 14, 4, 'Mandatory training · offshore', f'=SUM({R("mandatory", "earned")})/SUM({R("mandatory", "planned")})', PCT0,
     f'=TEXT(SUM({R("mandatory", "earned")}),"#,##0")&" / "&TEXT(SUM({R("mandatory", "planned")}),"#,##0")&" courses completed"')
proc_avg = lambda cat: f'AVERAGEIFS({R("procedures", "completion_pct")},{R("procedures", "category")},"{cat}")/100'
card(ov, r, 18, 4, 'Procedures · HSE | general operations', f'=TEXT({proc_avg("HSE")},"0%")&" | "&TEXT({proc_avg("General operations")},"0%")', 'General',
     f'=COUNTIFS({R("procedures", "issued")},"<>Yes",{R("procedures", "document")},"<>")&" procedures not yet issued"')
card(ov, r, 22, 4, 'PRs in progress', f'=SUM({R("pr_pipeline", "count")})', NUM,
     f'=SUMIFS({R("pr_pipeline", "count")},{R("pr_pipeline", "status")},"Late")&" late · "&SUMIFS({R("pr_pipeline", "count")},{R("pr_pipeline", "status")},"Near-late")'
     f'&" near-late · "&SUMIFS({R("packages_no_pr", "count")},{R("packages_no_pr", "dimension")},"status")&" packages without PR"')
ov.cell(r + 1, 6).font = F(size=22)

# manpower chart
ch = line_chart('Manpower ramp-up: plan vs actual (headcount)', 'People', w=30, h=9)
n = n_rows('manpower')
ch.add_data(data_ref('manpower', 'total_plan', n), titles_from_data=True)
ch.add_data(data_ref('manpower', 'total_actual', n), titles_from_data=True)
ch.set_categories(cat_ref('manpower', 'month', n))
style_series(ch.series[0], GREY, dashed=True); style_series(ch.series[1], BLUE)
ch.series[0].tx = SeriesLabel(v='Plan'); ch.series[1].tx = SeriesLabel(v='Actual')
ch.x_axis.tickLblSkip = 6
ov.add_chart(ch, 'B9')

# ORR curve
ch = line_chart('OPS teams ORR checklist vs plan (%)', '%', w=17, h=9)
n = n_rows('orr_curve')
ch.add_data(data_ref('orr_curve', 'plan_pct', n), titles_from_data=True)
ch.add_data(data_ref('orr_curve', 'actual_pct', n), titles_from_data=True)
ch.set_categories(cat_ref('orr_curve', 'month', n))
style_series(ch.series[0], GREY, dashed=True, marker=True); style_series(ch.series[1], BLUE, marker=True)
ch.series[0].tx = SeriesLabel(v='Plan'); ch.series[1].tx = SeriesLabel(v='Actual')
ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100
ov.add_chart(ch, 'R9')

# staffing
r = 28
section(ov, r, 2, 'Current staffing', 'Filled / required positions at the report date.')
table(ov, r + 2, 2, ['Group', '', 'Filled', 'Required', 'Filled %', ''])
body(ov, r + 3, 2, [lambda k: pick('staffing', 'group', k), lambda k: '', lambda k: pick('staffing', 'filled', k), lambda k: pick('staffing', 'required', k),
                    lambda k: f'=IF(D{r + 2 + k}="","",D{r + 2 + k}/E{r + 2 + k})', lambda k: ''], 5, [None, None, NUM, NUM, PCT0, None])
bar(ov, f'F{r + 3}:F{r + 7}')
red_below(ov, f'F{r + 3}:F{r + 7}', f'F{r + 3}')
ov.cell(r + 9, 2, f'="Vacant positions: "&IFERROR(INDEX({R("workstreams", "status")},MATCH("Vacant positions",{R("workstreams", "item")},0)),"—")').font = F(size=9, color=MUTED)

# ================================================================================================
# Manpower & Training
# ================================================================================================
pt = page('Manpower & Training', '2A9D8F')
header(pt, 'Mandatory training, CBTA development and assessment, vendor training schedule.')
section(pt, 4, 2, 'Mandatory training – PQPOC offshore staff', 'Courses completed vs required by the HSE training matrix. Red below 30%.')
table(pt, 6, 2, ['Team', '', 'Planned', 'Earned', 'Balance', 'Completed %'])
body(pt, 7, 2, [lambda k: pick('mandatory', 'team', k), lambda k: '', lambda k: pick('mandatory', 'planned', k), lambda k: pick('mandatory', 'earned', k),
                lambda k: pick('mandatory', 'balance', k), lambda k: f'=IF(D{6 + k}="","",E{6 + k}/D{6 + k})'], 10, [None, None, NUM, NUM, NUM, PCT0])
bar(pt, 'G7:G16'); red_below(pt, 'G7:G16', 'G7', 0.3)
pt.cell(17, 2, 'Total').font = F(bold=True)
pt['D17'] = '=SUM(D7:D16)'; pt['E17'] = '=SUM(E7:E16)'; pt['F17'] = '=SUM(F7:F16)'; pt['G17'] = '=IF(D17=0,"",E17/D17)'
for c in 'DEF': pt[f'{c}17'].number_format = NUM
pt['G17'].number_format = PCT0
for c in 'BCDEFG': pt[f'{c}17'].font = F(bold=True)

section(pt, 4, 10, 'CBTA by team: development and assessment', 'Modules developed and theory / OJT assessment progress.')
table(pt, 6, 10, ['Platform', 'Team', 'Modules', 'Completed', 'In dev.', 'In review', '% developed', '% theory', 'OJT modules', '% OJT'])
body(pt, 7, 10, [lambda k: pick('cbta', 'platform', k), lambda k: pick('cbta', 'team', k), lambda k: pick('cbta', 'modules', k), lambda k: pick('cbta', 'completed', k),
                 lambda k: pick('cbta', 'developing', k), lambda k: pick('cbta', 'review', k), lambda k: f'=IF(L{6 + k}="","",M{6 + k}/L{6 + k})',
                 lambda k: f'=IF(J{6 + k}="","",INDEX({R("cbta", "theory_pct")},{k})/100)', lambda k: pick('cbta', 'ojt_modules', k),
                 lambda k: f'=IF(J{6 + k}="","",INDEX({R("cbta", "ojt_pct")},{k})/100)'], 10, [None, None, NUM, NUM, NUM, NUM, PCT, PCT, NUM, PCT])
bar(pt, 'P7:P16'); bar(pt, 'Q7:Q16', '2A9D8F'); bar(pt, 'S7:S16', '7B5EA7')

n = n_rows('cbta')
ch = BarChart(); ch.type = 'bar'; ch.grouping = 'stacked'; ch.overlap = 100; ch.title = 'CBTA modules by status'; ch.height = 9; ch.width = 14
for k in range(1, 11):
    pt.cell(6 + k, 23, f'=IF(J{6 + k}="","",J{6 + k}&" "&K{6 + k})')
pt.column_dimensions['W'].hidden = True
for c, color, name in [(13, BLUE, 'Completed'), (15, AMBER, 'In review'), (14, GREY, 'In development')]:
    ch.add_data(Reference(pt, min_col=c, min_row=7, max_row=6 + n), titles_from_data=False)
    s = ch.series[-1]; s.tx = SeriesLabel(v=name); s.graphicalProperties.solidFill = color; s.graphicalProperties.line.solidFill = color
ch.set_categories(Reference(pt, min_col=23, min_row=7, max_row=6 + n)); ch.legend.position = 'b'; ch.y_axis.majorGridlines = None
ch.visible_cells_only = False
ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.scaling.orientation = 'maxMin'; ch.y_axis.crosses = 'max'
pt.add_chart(ch, 'B20')

section(pt, 20, 10, 'Vendor training schedule', 'Period = month (YYYY-MM) or year.')
table(pt, 22, 10, ['Period', '', 'Course', '', '', '', ''])
pt.merge_cells('L22:P22')
body(pt, 23, 10, [lambda k: pick('vendor_training', 'period', k), lambda k: '', lambda k: pick('vendor_training', 'course', k)], 30)
for k in range(30): pt.merge_cells(start_row=23 + k, start_column=12, end_row=23 + k, end_column=16)

# ================================================================================================
# ORR
# ================================================================================================
orr = page('ORR', '7B5EA7')
header(orr, 'ORR checklist: group totals, progress by OPS team and by checklist element.')
for i, g in enumerate(['TOTAL', 'PEOPLE', 'PLANT', 'PROCESS']):
    m = f'MATCH("{g}",{R("orr_groups", "group")},0)'
    card(orr, 4, 2 + i * 4, 4, f'="{g.title() if g != "TOTAL" else "Whole checklist"} · "&INDEX({R("orr_groups", "items")},{m})&" items"',
         f'=INDEX({R("orr_groups", "progress_pct")},{m})/100', PCT)

section(orr, 8, 2, '="Progress by OPS team · "&TEXT(DATEVALUE(' + orr_month + '&"-01"),"mmm yyyy")', 'Gap = actual – plan, in percentage points. Red when behind plan.')
table(orr, 10, 2, ['Team', '', 'Weight', 'Activities', 'Completed', 'Plan %', 'Actual %', 'Gap (pts)', 'Final plan %'])
body(orr, 11, 2, [lambda k: pick('orr_teams', 'team', k), lambda k: '', lambda k: f'=IF(B{10 + k}="","",INDEX({R("orr_teams", "weight_pct")},{k})/100)',
                  lambda k: pick('orr_teams', 'total_activities', k), lambda k: pick('orr_teams', 'completed', k),
                  lambda k: f'=IF(B{10 + k}="","",INDEX({R("orr_teams", "plan_pct")},{k})/100)', lambda k: f'=IF(B{10 + k}="","",INDEX({R("orr_teams", "actual_pct")},{k})/100)',
                  lambda k: f'=IF(B{10 + k}="","",(H{10 + k}-G{10 + k})*100)', lambda k: f'=IF(B{10 + k}="","",INDEX({R("orr_teams", "plan_end_pct")},{k})/100)'],
     12, [None, None, PCT, NUM, NUM, PCT, PCT, '+0.0;-0.0;0.0', PCT])
bar(orr, 'H11:H22')
orr.conditional_formatting.add('I11:I22', FormulaRule(formula=['AND(I11<>"",I11<0)'], font=Font(color=RED, bold=True)))
orr.conditional_formatting.add('I11:I22', FormulaRule(formula=['AND(I11<>"",I11>=0)'], font=Font(color=GREEN, bold=True)))

n = n_rows('orr_teams')
ch = BarChart(); ch.type = 'bar'; ch.title = 'Plan vs actual by team (%)'; ch.height = 9; ch.width = 15
for col, color, name in [('plan_pct', GREY, 'Plan %'), ('actual_pct', BLUE, 'Actual %')]:
    ch.add_data(data_ref('orr_teams', col, n), titles_from_data=True); ch.series[-1].graphicalProperties.solidFill = color; ch.series[-1].graphicalProperties.line.solidFill = color
    ch.series[-1].tx = SeriesLabel(v=name)
ch.set_categories(cat_ref('orr_teams', 'team', n)); ch.legend.position = 'b'; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100
ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.scaling.orientation = 'maxMin'; ch.y_axis.crosses = 'max'; ch.y_axis.majorGridlines = None
orr.add_chart(ch, 'L8')

section(orr, 25, 2, 'Checklist by element (People · Plant · Process)', 'Elements below 50% are shown in red.')
table(orr, 27, 2, ['Code', 'Element', '', '', '', '', 'Group', 'Items', 'Progress'])
body(orr, 28, 2, [lambda k: pick('orr_elements', 'code', k), lambda k: pick('orr_elements', 'element', k), *[lambda k: ''] * 4, lambda k: pick('orr_elements', 'group', k),
                  lambda k: pick('orr_elements', 'items', k), lambda k: f'=IF(B{27 + k}="","",INDEX({R("orr_elements", "progress_pct")},{k})/100)'],
     30, [None] * 7 + [NUM, PCT])
for k in range(30): orr.merge_cells(start_row=28 + k, start_column=3, end_row=28 + k, end_column=7)
bar(orr, 'J28:J57'); red_below(orr, 'J28:J57', 'J28')

# ================================================================================================
# Procedures
# ================================================================================================
pr = page('Procedures', '1F7A4D')
header(pr, 'Procedure completion. The full register is the "Quy trinh" data sheet; this page lists the ones still open.')
for i, (cat, label) in enumerate([('HSE', 'HSE procedures'), ('General operations', 'General operations procedures')]):
    cnt = f'COUNTIFS({R("procedures", "category")},"{cat}")'
    card(pr, 4, 2 + i * 5, 5, f'="{label} · "&{cnt}&" documents"', f'={proc_avg(cat)}', PCT0,
         f'=COUNTIFS({R("procedures", "category")},"{cat}",{R("procedures", "issued")},"<>Yes")&" not yet issued"')
section(pr, 8, 2, 'Open procedures (not yet issued)', 'In register order. Red: 0% · amber: in progress.')
table(pr, 10, 2, ['Code', 'Document', '', '', '', 'Owner', '', 'Due', 'Completion', 'Remark', '', '', '', '', ''])
cond = f'({R("procedures", "issued")}<>"Yes")*({R("procedures", "document")}<>"")'
for k in range(1, 41):
    rr = 10 + k
    pr.cell(rr, 26, nth_where('procedures', cond, k)).font = F(color='FFFFFF')  # helper index, hidden column Z
    for c, col in [(2, 'id'), (3, 'document'), (7, 'responsible'), (9, 'expected'), (10, 'completion_pct'), (11, 'remark')]:
        f = by_idx('procedures', col, f'$Z{rr}')
        if col == 'completion_pct': f = f'=IF($Z{rr}="","",INDEX({R("procedures", col)},$Z{rr})/100)'
        cell = pr.cell(rr, c, f); cell.font = F(); cell.fill = WHITE; cell.border = BOTTOM; cell.alignment = Alignment(vertical='top', wrap_text=True)
    pr.cell(rr, 10).number_format = PCT0
    pr.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
    pr.merge_cells(start_row=rr, start_column=7, end_row=rr, end_column=8)
    pr.merge_cells(start_row=rr, start_column=11, end_row=rr, end_column=16)
    pr.row_dimensions[rr].height = 28
pr.column_dimensions['Z'].hidden = True
pr.conditional_formatting.add('J11:J50', FormulaRule(formula=['AND(J11<>"",J11=0)'], font=Font(color=RED, bold=True)))
pr.conditional_formatting.add('J11:J50', FormulaRule(formula=['AND(J11<>"",J11>0,J11<1)'], font=Font(color='A8650A', bold=True)))

# ================================================================================================
# Contracts & Permits
# ================================================================================================
cp = page('Contracts & Permits', 'E0A526')
header(cp, 'Tendering pipeline, packages without PR, applications, logistics and permit register.')
section(cp, 4, 2, 'PRs in progress by stage', 'Count by schedule status. Stage names must match the "PR dang xu ly" data sheet.')
STAGES = ['1 CSP Preparation', '2 CSP Approval', '3 Bidding', '4 Evaluation', '5 TCEAR Approval', '6 Contract Signing']
table(cp, 6, 2, ['Stage', '', 'Late', 'Near-late', 'On time', 'Total'])
for k, s in enumerate(STAGES):
    rr = 7 + k
    cp.cell(rr, 2, s)
    for c, st in [(4, 'Late'), (5, 'Near-late'), (6, 'On time')]:
        cp.cell(rr, c, f'=SUMIFS({R("pr_pipeline", "count")},{R("pr_pipeline", "stage")},$B{rr},{R("pr_pipeline", "status")},"{st}")')
    cp.cell(rr, 7, f'=SUM(D{rr}:F{rr})')
    for c in range(2, 8):
        cp.cell(rr, c).font = F(); cp.cell(rr, c).fill = WHITE; cp.cell(rr, c).border = BOTTOM
cp.conditional_formatting.add('D7:D12', CellIsRule(operator='greaterThan', formula=['0'], font=Font(color=RED, bold=True)))
cp.conditional_formatting.add('E7:E12', CellIsRule(operator='greaterThan', formula=['0'], font=Font(color='A8650A', bold=True)))
ch = BarChart(); ch.type = 'bar'; ch.grouping = 'stacked'; ch.overlap = 100; ch.title = 'PRs by stage'; ch.height = 7; ch.width = 13
for c, color in [(4, RED), (5, AMBER), (6, BLUE)]:
    ch.add_data(Reference(cp, min_col=c, min_row=6, max_row=12), titles_from_data=True)
    ch.series[-1].graphicalProperties.solidFill = color; ch.series[-1].graphicalProperties.line.solidFill = color
ch.set_categories(Reference(cp, min_col=2, min_row=7, max_row=12)); ch.legend.position = 'b'
ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.scaling.orientation = 'maxMin'; ch.y_axis.crosses = 'max'; ch.y_axis.majorGridlines = None
cp.add_chart(ch, 'B15')

section(cp, 4, 10, 'Packages without PR')
rr0 = 6
for dim, label in [('status', 'Procurement status'), ('priority', 'Priority'), ('scr_timing', 'Expected SCR')]:
    table(cp, rr0, 10, [label, '', 'Packages'])
    cond = f'{R("packages_no_pr", "dimension")}="{dim}"'
    for k in range(1, 6):
        rr = rr0 + k
        cp.cell(rr, 25, nth_where('packages_no_pr', cond, k)).font = F(color='FFFFFF')
        for c, col in [(10, 'category'), (12, 'count')]:
            cell = cp.cell(rr, c, by_idx('packages_no_pr', col, f'$Y{rr}')); cell.font = F(); cell.fill = WHITE; cell.border = BOTTOM
        cp.cell(rr, 11).fill = WHITE; cp.cell(rr, 11).border = BOTTOM
    cp.conditional_formatting.add(f'L{rr0 + 1}:L{rr0 + 5}', DataBarRule(start_type='num', start_value=0, end_type='max', color=BLUE, showValue=True))
    rr0 += 7
cp.column_dimensions['Y'].hidden = True

section(cp, 4, 14, 'CMMS (OA) – actual vs plan')
table(cp, 6, 14, ['Package', 'Actual', 'Plan', 'Gap (pts)'])
body(cp, 7, 14, [lambda k: pick('cmms', 'package', k), lambda k: f'=IF(N{6 + k}="","",INDEX({R("cmms", "actual_pct")},{k})/100)',
                 lambda k: f'=IF(N{6 + k}="","",INDEX({R("cmms", "plan_pct")},{k})/100)', lambda k: f'=IF(N{6 + k}="","",(O{6 + k}-P{6 + k})*100)'],
     4, [None, PCT, PCT, '+0.0;-0.0;0.0'])
bar(cp, 'O7:O10')

def status_list(ws, r, c, title, areas, n=12, helper=26):
    section(ws, r, c, title)
    table(ws, r + 2, c, ['Item', '', 'Status', '', '', '', '', ''])
    cond = '+'.join(f'({R("workstreams", "area")}="{a}")' for a in areas)
    for k in range(1, n + 1):
        rr = r + 2 + k
        ws.cell(rr, helper, nth_where('workstreams', f'({cond})>0', k)).font = F(color='FFFFFF')
        for cc, col, span in [(c, 'item', 2), (c + 2, 'status', 6)]:
            cell = ws.cell(rr, cc, by_idx('workstreams', col, f'${L(helper)}{rr}')); cell.font = F(bold=(col == 'item')); cell.alignment = Alignment(vertical='top', wrap_text=True)
            ws.merge_cells(start_row=rr, start_column=cc, end_row=rr, end_column=cc + span - 1)
            for x in range(cc, cc + span): ws.cell(rr, x).fill = WHITE; ws.cell(rr, x).border = BOTTOM
        ws.row_dimensions[rr].height = 30
    ws.column_dimensions[L(helper)].hidden = True

status_list(cp, 28, 2, 'Applications', ['Applications'], 8)
status_list(cp, 28, 11, 'Logistics, services & coordination', ['Logistics & services', 'Coordination'], 10, helper=24)

section(cp, 42, 2, 'Offshore platform permit register', 'Red: waiting · amber: temporary · green: issued.')
table(cp, 44, 2, ['Permit', '', '', '', 'Applies to', '', 'Authority', '', 'Issued', 'Expires', 'Renew by', 'Status', '', ''])
for k in range(1, 21):
    rr = 44 + k
    for c, col, span in [(2, 'permit', 4), (6, 'facility', 2), (8, 'authority', 2), (10, 'issue_date', 1), (11, 'expiry_date', 1), (12, 'renewal_date', 1), (13, 'status', 3)]:
        cell = cp.cell(rr, c, pick('permits', col, k)); cell.font = F(); cell.alignment = Alignment(vertical='top', wrap_text=True)
        if span > 1: cp.merge_cells(start_row=rr, start_column=c, end_row=rr, end_column=c + span - 1)
        for x in range(c, c + span): cp.cell(rr, x).fill = WHITE; cp.cell(rr, x).border = BOTTOM
    cp.row_dimensions[rr].height = 28
cp.conditional_formatting.add('M45:M64', FormulaRule(formula=['ISNUMBER(SEARCH("Waiting",M45))'], font=Font(color=RED, bold=True)))
cp.conditional_formatting.add('M45:M64', FormulaRule(formula=['M45="Temporary"'], font=Font(color='A8650A', bold=True)))
cp.conditional_formatting.add('M45:M64', FormulaRule(formula=['M45="Issued"'], font=Font(color=GREEN, bold=True)))

# ================================================================================================
# Commissioning
# ================================================================================================
cm = page('Commissioning', 'D1495B')
header(cm, 'CPP pre-commissioning plan by week. Green: done (actual finish entered) · red: in its window on the report date · blue: planned.')
cm.column_dimensions['B'].width = 46
for c in range(3, 6): cm.column_dimensions[L(c)].width = 10
WEEKS, W0 = 22, 6
cm['C4'] = 'Report date'; cm['D4'] = "=DATEVALUE('Thong tin'!$A$2)"; cm['D4'].number_format = 'd mmm yyyy'
cm['C4'].font = F(size=9, color=MUTED); cm['D4'].font = F(bold=True)
starts = R('precom_plan', 'plan_start')
cm['C5'] = 'First week'; cm['C5'].font = F(size=9, color=MUTED)
first_start = f'IF(COUNTIF({starts},"?*")=0,D4,MIN(IF({starts}<>"",DATEVALUE({starts}))))'
cm['D5'] = f'={first_start}-WEEKDAY({first_start},2)+1'
cm['D5'].number_format = 'd mmm yyyy'; cm['D5'].font = F(bold=True)
table(cm, 6, 2, ['Activity', 'Plan start', 'Plan finish', 'Actual finish'])
for w in range(WEEKS):
    c = W0 + w
    cm.column_dimensions[L(c)].width = 4.2
    cell = cm.cell(6, c, '=$D$5' if w == 0 else f'={L(c - 1)}6+7')
    cell.number_format = 'd mmm'; cell.font = F(size=8, color=MUTED); cell.fill = HEADF
    cell.alignment = Alignment(text_rotation=90, horizontal='center')
cm.row_dimensions[6].height = 48
NG = 30
for k in range(1, NG + 1):
    rr = 6 + k
    cm.cell(rr, 2, pick('precom_plan', 'activity', k))
    for c, col in [(3, 'plan_start'), (4, 'plan_finish'), (5, 'actual_finish')]:
        cm.cell(rr, c, f'=IF(INDEX({R("precom_plan", col)},{k})="","",DATEVALUE(INDEX({R("precom_plan", col)},{k})))').number_format = 'd mmm yy'
    for c in range(2, W0 + WEEKS):
        cell = cm.cell(rr, c); cell.font = F(size=9); cell.fill = WHITE; cell.border = BOTTOM
        cell.alignment = Alignment(vertical='center', wrap_text=(c == 2))
    cm.row_dimensions[rr].height = 26
grid = f'{L(W0)}7:{L(W0 + WEEKS - 1)}{6 + NG}'
first = f'{L(W0)}7'
inwin = f'AND($C7<>"",{L(W0)}$6<=$D7,{L(W0)}$6+6>=$C7)'
cm.conditional_formatting.add(grid, FormulaRule(formula=[f'AND({inwin},$E7<>"")'], fill=PatternFill('solid', fgColor='5CC28D'), stopIfTrue=True))
cm.conditional_formatting.add(grid, FormulaRule(formula=[f'AND({inwin},$C7<=$D$4,$D7>=$D$4)'], fill=PatternFill('solid', fgColor=RED), stopIfTrue=True))
cm.conditional_formatting.add(grid, FormulaRule(formula=[inwin], fill=PatternFill('solid', fgColor='9DB9D6'), stopIfTrue=True))
cm.conditional_formatting.add(grid, FormulaRule(formula=[f'AND({L(W0)}$6<=$D$4,{L(W0)}$6+6>=$D$4)'], border=Border(left=Side(style='medium', color=NAVY))))

r0 = 8 + NG
section(cm, r0, 2, 'Preservation, installation & commissioning')
table(cm, r0 + 2, 2, ['Item', 'Status', '', ''])
cond = f'{R("workstreams", "area")}="Commissioning"'
for k in range(1, 6):
    rr = r0 + 2 + k
    cm.cell(rr, 26, nth_where('workstreams', cond, k)).font = F(color='FFFFFF')
    cm.cell(rr, 2, by_idx('workstreams', 'item', f'$Z{rr}')).font = F(bold=True)
    cm.cell(rr, 3, by_idx('workstreams', 'status', f'$Z{rr}')).font = F()
    cm.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=14)
cm.column_dimensions['Z'].hidden = True

r1 = r0 + 10
section(cm, r1, 2, 'CPP start-up sequence', 'Summarised from the CPP start-up block diagram.')
table(cm, r1 + 2, 2, ['Step · phase', 'Action'])
for k in range(1, 21):
    rr = r1 + 2 + k
    cm.cell(rr, 2, f'=IF(INDEX({R("startup_sequence", "step")},{k})="","",INDEX({R("startup_sequence", "step")},{k})&". "&INDEX({R("startup_sequence", "phase")},{k}))').font = F(bold=True)
    cm.cell(rr, 3, pick('startup_sequence', 'action', k)).font = F()
    cm.cell(rr, 3).alignment = Alignment(wrap_text=True, vertical='top')
    cm.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=W0 + WEEKS - 1)
    cm.row_dimensions[rr].height = 30

# formulas that filter a range must be array formulas so every Excel version evaluates the range
from openpyxl.worksheet.formula import ArrayFormula
for s in (ov, pt, orr, pr, cp, cm):
    for row in s.iter_rows():
        for c in row:
            if isinstance(c.value, str) and ('SMALL(IF(' in c.value or 'MIN(IF(' in c.value):
                c.value = ArrayFormula(c.coordinate, c.value)

# ---- order: dashboards first, then guide, then data ------------------------------------------
dash = ['Overview', 'Manpower & Training', 'ORR', 'Procedures', 'Contracts & Permits', 'Commissioning']
wb._sheets = [wb[s] for s in dash] + [s for s in wb._sheets if s.title not in dash]
wb.active = 0
for s in wb.worksheets: s.sheet_view.tabSelected = (s.title == 'Overview')
for s in (ov, pt, orr, pr, cp, cm):  # print / PDF: one page wide, landscape
    s.page_setup.orientation = 'landscape'; s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT, [s.title for s in wb.worksheets][:8])
