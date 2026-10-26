"""Build OPS_Dashboard_Data.xlsx: one Excel Table per dashboard dataset + a guide sheet.

usage: excel.py <data dir> <out.xlsx>
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

D, OUT = sys.argv[1:3]

# table id -> (sheet name, what it feeds, {column: Vietnamese meaning})
SPEC = {
 'meta': ('Thong tin', 'Ngày báo cáo và tiêu đề (1 dòng duy nhất)', {
   'report_date': 'Ngày báo cáo, dạng YYYY-MM-DD', 'report_title': 'Tên báo cáo', 'prepared_by': 'Đơn vị lập'}),
 'manpower': ('Nhan su theo thang', 'Biểu đồ tăng cường nhân sự; KPI nhân sự lấy tháng cuối cùng có số thực tế', {
   'month': 'Tháng, dạng YYYY-MM', 'onshore_plan': 'Onshore kế hoạch', 'onshore_actual': 'Onshore thực tế (để trống nếu chưa tới)',
   'offshore_plan': 'Offshore kế hoạch', 'offshore_actual': 'Offshore thực tế', 'total_plan': 'Tổng kế hoạch', 'total_actual': 'Tổng thực tế'}),
 'staffing': ('Dinh bien', 'Thẻ "Định biên hiện tại"', {'group': 'Nhóm', 'filled': 'Đã có', 'required': 'Cần'}),
 'cbta': ('CBTA', 'Tab Nhân sự & đào tạo: CBTA', {
   'platform': 'CPP hoặc WHP', 'team': 'Team', 'modules': 'Tổng module', 'developing': 'Đang biên soạn', 'review': 'Đang soát xét',
   'completed': 'Đã hoàn thành', 'theory_pct': '% đánh giá lý thuyết (0–100)', 'ojt_modules': 'Số module OJT', 'ojt_pct': '% đánh giá OJT (0–100)'}),
 'mandatory': ('Dao tao bat buoc', 'Đào tạo bắt buộc offshore', {
   'team': 'Team', 'planned': 'Số khoá kế hoạch', 'earned': 'Số khoá đã hoàn thành', 'balance': 'Còn lại'}),
 'vendor_training': ('Dao tao vendor', 'Lịch đào tạo vendor', {'course': 'Tên khoá', 'period': 'Tháng YYYY-MM, hoặc chỉ năm YYYY'}),
 'orr_teams': ('ORR theo team', 'ORR theo team OPS (tháng hiện tại)', {
   'team': 'Team', 'weight_pct': 'Trọng số %', 'total_activities': 'Tổng hoạt động', 'completed': 'Đã hoàn thành',
   'plan_pct': '% kế hoạch tháng hiện tại', 'actual_pct': '% thực tế tháng hiện tại', 'plan_end_pct': '% kế hoạch cuối kỳ'}),
 'orr_curve': ('ORR theo thang', 'Đường ORR kế hoạch/thực tế; KPI ORR lấy tháng cuối cùng có thực tế', {
   'month': 'Tháng YYYY-MM', 'plan_pct': '% kế hoạch tổng', 'actual_pct': '% thực tế tổng (để trống nếu chưa tới)'}),
 'orr_elements': ('ORR hang muc', 'Bảng checklist theo hạng mục', {
   'group': 'PEOPLE / PLANT / PROCESS', 'code': 'Mã (không trùng)', 'element': 'Tên hạng mục', 'items': 'Số đầu mục', 'progress_pct': '% tiến độ (0–100)'}),
 'orr_groups': ('ORR tong nhom', 'Ô tổng ORR theo nhóm', {
   'group': 'PEOPLE / PLANT / PROCESS / TOTAL (giữ nguyên tên)', 'items': 'Số đầu mục', 'progress_pct': '% tiến độ (0–100)'}),
 'procedures': ('Quy trinh', 'Tab Quy trình', {
   'id': 'Mã (không trùng, ví dụ HSE-41)', 'category': '"HSE" hoặc "General operations" (giữ đúng chữ)', 'document': 'Tên tài liệu',
   'level': 'Company / Department', 'issued': '"Yes" nếu đã ban hành, "No" nếu chưa', 'responsible': 'Người phụ trách',
   'expected': 'Hạn, ví dụ Dec-26', 'completion_pct': '% hoàn thành (0–100)', 'remark': 'Ghi chú'}),
 'packages_no_pr': ('Goi chua co PR', 'Gói chưa có PR', {
   'dimension': 'status / priority / scr_timing (giữ nguyên)', 'category': 'Nhóm', 'count': 'Số gói'}),
 'pr_pipeline': ('PR dang xu ly', 'PR theo bước', {
   'stage': 'Bước (giữ số thứ tự ở đầu)', 'status': '"Late", "Near-late" hoặc "On time" (giữ đúng chữ)', 'count': 'Số PR'}),
 'cmms': ('CMMS', 'CMMS (OA)', {'package': 'Gói', 'actual_pct': '% thực tế', 'plan_pct': '% kế hoạch'}),
 'workstreams': ('Trang thai cong viec', 'Danh sách ứng dụng, logistics, điều phối, commissioning, vị trí trống', {
   'area': '"Manpower", "Applications", "Logistics & services", "Coordination" hoặc "Commissioning" (giữ đúng chữ)',
   'item': 'Hạng mục ("Vacant positions" là dòng vị trí còn trống)', 'status': 'Trạng thái / bước tiếp theo'}),
 'permits': ('Giay phep', 'Đăng ký giấy phép', {
   'permit': 'Tên giấy phép (không trùng)', 'facility': 'Áp dụng cho', 'authority': 'Cơ quan', 'issue_date': 'Ngày cấp YYYY-MM-DD',
   'expiry_date': 'Ngày hết hạn YYYY-MM-DD', 'renewal_date': 'Gia hạn trước YYYY-MM-DD', 'status': '"Issued", "Temporary", "Waiting…" hoặc ghi chú'}),
 'precom_plan': ('Pre-com', 'Biểu đồ tiến độ pre-commissioning', {
   'activity': 'Tên hoạt động', 'plan_start': 'Bắt đầu KH YYYY-MM-DD', 'plan_finish': 'Kết thúc KH YYYY-MM-DD',
   'actual_start': 'Bắt đầu thực tế YYYY-MM-DD', 'actual_finish': 'Kết thúc thực tế YYYY-MM-DD (có ngày → thanh màu xanh)'}),
 'startup_sequence': ('Trinh tu khoi dong', 'Trình tự khởi động CPP', {'step': 'Số thứ tự', 'phase': 'Tên bước', 'action': 'Nội dung'}),
}
DATE_COLS = {'month', 'period', 'report_date', 'issue_date', 'expiry_date', 'renewal_date', 'plan_start', 'plan_finish', 'actual_start', 'actual_finish', 'expected'}

wb = Workbook()
g = wb.active
g.title = 'Huong dan'
HEAD = Font(bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='1F4E79')
rows = [
 ['OPS DASHBOARD – FILE SỐ LIỆU NGUỒN'], [],
 ['Cách dùng'],
 ['1. Mỗi sheet là một bảng (Excel Table). Chỉ sửa số liệu trong bảng; thêm dòng bằng cách gõ ngay dưới dòng cuối của bảng (bảng tự mở rộng).'],
 ['2. KHÔNG đổi tên bảng (Table Design → Table Name), KHÔNG đổi tên cột ở dòng tiêu đề. Dashboard đọc theo đúng các tên này.'],
 ['3. Ngày tháng gõ dạng chữ: tháng YYYY-MM (2026-11), ngày YYYY-MM-DD (2026-11-30). Các cột ngày đã định dạng Text sẵn.'],
 ['4. Phần trăm gõ số 0–100 (ví dụ 70.5), không gõ dấu %. Ô chưa có số để trống.'],
 ['5. MỖI KỲ BÁO CÁO: (a) sheet "Thong tin" → sửa report_date thành ngày báo cáo mới; (b) nhập số vào các bảng; (c) xem lại sheet Overview; (d) bấm nút / chạy script "Day so lieu len dashboard". Trang web cập nhật sau 1–2 phút.'],
 ['   Mỗi report_date được lưu thành một bản báo cáo riêng, xem lại trên web ở ô "Report". Đẩy lại cùng report_date sẽ ghi đè bản của ngày đó (dùng để sửa số).'],
 ['6. Các sheet dashboard (Overview … Commissioning) ở đầu file tự tính lại ngay khi sửa số; không cần sửa gì trên các sheet đó.'],
 [],
 ['Bảng', 'Sheet', 'Hiển thị ở', 'Cột', 'Ý nghĩa'],
]
for r in rows: g.append(r)
hdr = g.max_row
for c in g[hdr]: c.font = HEAD; c.fill = FILL
g['A1'].font = Font(bold=True, size=14); g['A3'].font = Font(bold=True)
for tid, (sheet, where, cols) in SPEC.items():
    first = True
    for col, desc in cols.items():
        g.append([tid if first else '', sheet if first else '', where if first else '', col, desc]); first = False
for col, w in zip('ABCDE', (18, 22, 46, 18, 70)): g.column_dimensions[col].width = w
for row in g.iter_rows(min_row=hdr + 1):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')

for tid, (sheet, where, cols) in SPEC.items():
    rows = json.load(open(os.path.join(D, tid + '.json')))
    ws = wb.create_sheet(sheet)
    names = list(cols)
    ws.append(names)
    for r in rows:
        ws.append([r.get(c) for c in names])
    n = len(rows)
    for j, c in enumerate(names, 1):
        L = get_column_letter(j)
        if c in DATE_COLS:
            for i in range(2, n + 202):  # pre-format room for new rows as text too
                ws[f'{L}{i}'].number_format = '@'
        width = max(len(c), *(len(str(r.get(c) or '')) for r in rows)) if rows else len(c)
        ws.column_dimensions[L].width = min(60, max(10, width + 2))
    for c in ws[1]: c.comment = None
    t = Table(displayName=tid, ref=f'A1:{get_column_letter(len(names))}{max(2, n + 1)}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = 'A2'
wb.save(OUT)
print('saved', OUT)
